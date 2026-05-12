from __future__ import annotations

import html
import json
import re
import textwrap
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "storyboard_output" / "ep5_director_image2"
PROMPT_DIR = OUT_DIR / "prompts"
BLOCKING_DIR = OUT_DIR / "blocking"
IMAGE_DIR = OUT_DIR / "image2_results"
WORKBOOK = ROOT / "十年春风分镜头表_第15集已补.xlsx"

REF_A_PATH = Path(r"D:\wechat\xwechat_files\ldmsapp_46cb\temp\RWTemp\2026-05\e0b5ff5ed172ccd032d4185083521c8a\b07d6083e4cdcd9f8efcec7948fd9ff4.jpg")
REF_B_PATH = Path(r"D:\wechat\xwechat_files\ldmsapp_46cb\temp\RWTemp\2026-05\e0b5ff5ed172ccd032d4185083521c8a\ac4e54248cf99e0943d440d7a7c2eddc.jpg")
REF_A_LABEL = "参考图2：回廊正向视角，外侧园林在左，红墙在右"
REF_B_LABEL = "参考图3：回廊反向视角，红墙在左，外侧园林在右"


@dataclass
class Shot:
    sid: str
    duration: int
    device: str
    size: str
    angle: str
    move: str
    blocking: str
    lighting: str
    content: str
    transition: str
    image_prompt: str


def clean(value: object) -> str:
    text = str(value or "").strip()
    return re.sub(r"^[^：]+：", "", text).strip()


def seconds(value: object) -> int:
    match = re.search(r"\d+", str(value or ""))
    return int(match.group()) if match else 0


def esc(value: object) -> str:
    return html.escape(str(value or ""), quote=True)


def read_shots() -> list[Shot]:
    wb = load_workbook(WORKBOOK, data_only=True)
    ws = wb["第五集"]
    shots: list[Shot] = []
    for row in range(3, ws.max_row + 1):
        values = [ws.cell(row, col).value for col in range(1, 12)]
        if not values[0]:
            continue
        sid = clean(values[0])
        shot = Shot(
            sid=sid,
            duration=seconds(values[1]),
            device=clean(values[2]),
            size=clean(values[3]),
            angle=clean(values[4]),
            move=clean(values[5]),
            blocking=clean(values[6]),
            lighting=clean(values[7]),
            content=str(values[8] or "").strip(),
            transition=clean(values[9]),
            image_prompt="",
        )
        if sid == "5-19":
            shot.blocking = "谢临舟走出两步后停下，半侧回头看向苏明舒，石墩停在他身后；苏明舒站在回廊暖光处抬眼回应。"
            shot.lighting = "谢临舟侧脸用冷调轮廓光，回廊纵深压暗；远处苏明舒保留暖光虚化，突出回头补刀的钩子。"
            shot.content = "谢临舟：机关不错。苏明舒挑眉。谢临舟补一句：就是火候差点。[环境声：脚步声停住、风过回廊]"
            shot.transition = "正反打切换"
        shot.image_prompt = make_image_prompt(shot)
        shots.append(shot)
    return shots


def make_image_prompt(shot: Shot) -> str:
    lens_note = "参考图2的回廊方向" if shot.sid in {"5-1", "5-2", "5-4", "5-6", "5-10", "5-13", "5-16", "5-24"} else "参考图3的反向回廊方向"
    if shot.sid in {"5-22", "5-23"}:
        lens_note = "沿用回廊反向视角的冷暗背景，重点给袖中锁灵玉和谢临舟侧脸"
    subject = re.sub(r"\[[^\]]+\]", "", shot.content)
    return (
        f"Codex imagegen，9:16竖幅，黑白铅笔导演分镜草图，古装短剧《十年春风与君归》第5集镜头{shot.sid}。"
        f"场景是御花园木结构长回廊，红墙、木柱、漏窗、石板地、外侧园林树影，{lens_note}。"
        f"镜头景别：{shot.size}；拍摄视角：{shot.angle}；运镜意图：{shot.move}。"
        f"人物调度：{shot.blocking}。"
        f"画面动作：{subject}。"
        "角色造型：苏明舒为温婉聪慧古装贵女，浅色衣裙；谢临舟为冷峻锦衣卫，玄色衣袍；青枝为侍女；石墩为锦衣卫随从。"
        "画面只要分镜草图线稿和灰阶明暗，不要彩色，不要字幕，不要对白文字，不要标题，不要水印，不要边框文字。"
        "构图必须是单张独立分镜画面，不要拼图，不要漫画格。人物脸部清晰，古装真实，电影感，导演故事板草图。"
    )


def batch_prompt(shots: list[Shot]) -> str:
    lines = [
        "请使用 Codex 内置 imagegen 生成古装短剧导演分镜草图。",
        "输出要求：每个镜头生成一张独立图片，共生成本批次列出的所有图片；全部为9:16竖幅；黑白铅笔故事板草图；无字幕、无标题、无对白文字、无水印、不要拼图。",
        f"{REF_A_LABEL}",
        f"{REF_B_LABEL}",
        "两张参考图是同一御花园回廊的正反视角，所有镜头都要保持这个木廊、红墙、漏窗、石板地、园林树影的空间连续性。",
        "",
    ]
    for shot in shots:
        lines.append(f"【{shot.sid}】{shot.image_prompt}")
    return "\n".join(lines)


COLORS = {
    "苏明舒": "#d89b65",
    "谢临舟": "#586a83",
    "青枝": "#78a06f",
    "石墩": "#8b7a68",
    "锁灵玉": "#82c9c2",
}


def positions(sid: str) -> dict[str, tuple[int, int]]:
    n = int(sid.split("-")[1])
    if n in {1, 16, 17, 19, 21, 22, 23, 24}:
        return {"苏明舒": (124, 226), "青枝": (82, 262), "谢临舟": (232, 122), "石墩": (280, 98)}
    if n in {2, 5, 7, 9, 14, 18, 20}:
        return {"苏明舒": (146, 204), "青枝": (92, 248), "谢临舟": (246, 128), "石墩": (288, 106)}
    if n in {3, 8, 12}:
        return {"谢临舟": (174, 132), "石墩": (246, 112), "苏明舒": (98, 232)}
    if n in {13, 15}:
        return {"苏明舒": (152, 210), "青枝": (92, 250), "谢临舟": (202, 172), "石墩": (268, 132)}
    return {"苏明舒": (142, 220), "青枝": (88, 252), "谢临舟": (232, 142), "石墩": (286, 116)}


def camera(sid: str) -> tuple[int, int, int, int]:
    n = int(sid.split("-")[1])
    if n in {1, 16, 24}:
        return (78, 312, 202, 156)
    if n in {3, 8, 12, 17, 19, 21, 23}:
        return (284, 306, 186, 148)
    if n in {7, 14, 18, 20}:
        return (134, 310, 146, 202)
    return (178, 310, 188, 178)


def blocking_svg(shot: Shot) -> str:
    cx, cy, tx, ty = camera(shot.sid)
    parts = [
        '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 360 340">',
        '<defs><marker id="arr" markerWidth="12" markerHeight="12" refX="10" refY="6" orient="auto"><path d="M2,2 L10,6 L2,10 Z" fill="#356d9b"/></marker></defs>',
        '<rect width="360" height="340" rx="12" fill="#fffdf8"/>',
        '<rect x="36" y="38" width="288" height="86" rx="8" fill="#eee6dc" stroke="#d7c9bb"/>',
        '<text x="180" y="84" text-anchor="middle" class="muted">御花园回廊 / 纵深轴线</text>',
        '<path d="M56 126 L304 126 L324 292 L36 292 Z" fill="#f2eadf" stroke="#d7c9bb"/>',
        '<line x1="94" y1="132" x2="60" y2="286" stroke="#b79573" stroke-width="6"/>',
        '<line x1="266" y1="132" x2="300" y2="286" stroke="#b79573" stroke-width="6"/>',
        f'<path d="M {cx} {cy} L {tx - 70} {ty - 48} L {tx + 70} {ty - 48} Z" fill="#7aa0c4" opacity=".15" stroke="#356d9b"/>',
        f'<line x1="{cx}" y1="{cy}" x2="{tx}" y2="{ty}" stroke="#356d9b" stroke-width="4" marker-end="url(#arr)"/>',
        f'<polygon points="{cx - 18},{cy + 13} {cx + 18},{cy + 13} {cx},{cy - 22}" fill="#356d9b"/>',
        f'<text x="{cx}" y="{cy + 36}" text-anchor="middle" class="cam">CAM</text>',
    ]
    for name, (x, y) in positions(shot.sid).items():
        color = COLORS[name]
        parts.append(f'<circle cx="{x}" cy="{y}" r="20" fill="{color}" stroke="#222" stroke-width="2"/>')
        parts.append(f'<text x="{x}" y="{y + 40}" text-anchor="middle" class="name">{name}</text>')
    parts.append("</svg>")
    return "".join(parts)


def write_outputs(shots: list[Shot]) -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    PROMPT_DIR.mkdir(parents=True, exist_ok=True)
    BLOCKING_DIR.mkdir(parents=True, exist_ok=True)
    IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    for i in range(0, len(shots), 6):
        batch = shots[i:i + 6]
        (PROMPT_DIR / f"batch_{i // 6 + 1:02d}.txt").write_text(batch_prompt(batch), encoding="utf-8")
    for shot in shots:
        (PROMPT_DIR / f"{shot.sid}.txt").write_text(shot.image_prompt, encoding="utf-8")
        svg = blocking_svg(shot)
        svg = svg.replace(
            "<svg ",
            '<svg class="blocking-svg" ',
            1,
        )
        (BLOCKING_DIR / f"{shot.sid}.svg").write_text(svg, encoding="utf-8")

    html_doc = make_html(shots)
    (OUT_DIR / "ep5_director_storyboard_image2.html").write_text(html_doc, encoding="utf-8")
    md_doc = make_md(shots)
    (OUT_DIR / "ep5_director_storyboard_image2.md").write_text(md_doc, encoding="utf-8")


def image_tag(shot: Shot) -> str:
    candidates = sorted(IMAGE_DIR.glob(f"*{shot.sid.replace('-', '_')}*.*")) + sorted(IMAGE_DIR.glob(f"*{shot.sid}*.*"))
    if candidates:
        rel = candidates[0].relative_to(OUT_DIR).as_posix()
        return f'<img class="sketch" src="{esc(rel)}" alt="{esc(shot.sid)} Image 2.0 sketch">'
    return (
        '<div class="sketch placeholder">'
        f'<div>{esc(shot.sid)}</div><small>Image 2.0草图生成后会替换到这里</small>'
        '</div>'
    )


def make_html(shots: list[Shot]) -> str:
    cards = []
    for shot in shots:
        cards.append(
            f"""
<section class="shot-card">
  <header><b>镜头号：{esc(shot.sid)}</b><span>{esc(shot.duration)}秒</span></header>
  <div class="grid">
    <div>
      {image_tag(shot)}
      <p class="note">9:16 Image 2.0草图；参考回廊正反视角；图内不放文字。</p>
    </div>
    <div class="details">
      <p><strong>设备/镜头</strong>{esc(shot.device)}</p>
      <p><strong>景别</strong>{esc(shot.size)}</p>
      <p><strong>拍摄视角</strong>{esc(shot.angle)}</p>
      <p><strong>运镜</strong>{esc(shot.move)}</p>
      <p><strong>站位调度</strong>{esc(shot.blocking)}</p>
      <p><strong>布光</strong>{esc(shot.lighting)}</p>
      <p><strong>画面+台词+声音</strong>{esc(shot.content)}</p>
      <p><strong>转场</strong>{esc(shot.transition)}</p>
    </div>
    <div class="blocking">{(BLOCKING_DIR / f"{shot.sid}.svg").read_text(encoding="utf-8")}</div>
  </div>
</section>
"""
        )
    return f"""<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>第5集｜导演分镜表 + Image 2.0草图 + 人物站位</title>
<style>
body{{margin:0;background:#e6dfd4;color:#1f1b17;font-family:"Microsoft YaHei","Noto Sans SC",Arial,sans-serif}}
.page{{max-width:1500px;margin:0 auto;padding:18px 22px 40px}}
.hero{{background:#fffaf1;border:1px solid #d7cbbc;border-radius:8px;padding:18px 22px;margin-bottom:18px}}
h1{{font-size:26px;margin:0 0 10px}} .meta{{display:grid;grid-template-columns:repeat(4,1fr);gap:10px;font-size:14px;color:#5f564d}}
.overview{{margin:16px 0 0;max-width:620px}} .overview img{{width:100%;display:block;border:1px solid #d6cabb;background:#fff}} .overview figcaption{{font-size:13px;color:#6b6258;margin-top:6px}}
.shot-card{{background:#fffdf8;border:1px solid #d6cabb;border-radius:8px;margin:18px 0;padding:12px}}
.shot-card header{{display:flex;justify-content:space-between;align-items:center;font-size:18px;color:#7a481a;margin-bottom:10px}}
.shot-card header span{{background:#efe3d1;border-radius:18px;padding:3px 10px;color:#7b6042;font-size:14px}}
.grid{{display:grid;grid-template-columns:minmax(280px,390px) minmax(380px,1fr) 300px;gap:16px;align-items:start}}
.sketch{{width:100%;aspect-ratio:9/16;object-fit:cover;border:1px solid #333;background:#f2eee8;display:flex;align-items:center;justify-content:center;text-align:center;color:#7a6042;font-weight:700}}
.placeholder small{{display:block;font-size:13px;margin-top:8px;font-weight:400}}
.note{{font-size:12px;color:#7d756d;margin:6px 0 0}}
.details p{{border-bottom:1px dashed #dfd3c4;margin:0;padding:8px 0;line-height:1.65;font-size:15px}}
.details strong{{display:block;color:#8a551e;margin-bottom:2px}}
.blocking{{border:1px solid #d6cabb;border-radius:8px;overflow:hidden;background:#fff}}
.blocking-svg{{width:100%;display:block}}
.muted{{font-size:14px;fill:#6b6258;font-weight:700}} .name{{font-size:13px;fill:#1f1b17;font-weight:700}} .cam{{font-size:12px;fill:#356d9b;font-weight:700}}
@media(max-width:980px){{.grid{{grid-template-columns:1fr}}.meta{{grid-template-columns:1fr 1fr}}}}
</style>
</head>
<body><main class="page">
<section class="hero">
<h1>第5集《活阎罗谢临舟（下）》｜导演分镜表 + Image 2.0草图 + 人物关系位置图</h1>
<div class="meta"><div>镜头数：{len(shots)}</div><div>总时长：{sum(s.duration for s in shots)}秒</div><div>草图比例：9:16竖幅</div><div>场景参考：回廊正反视角图</div></div>
<figure class="overview"><img src="image2_results/ep5_24panel_overview_codex_imagegen.png" alt="第5集24镜Codex imagegen草图总览"><figcaption>Codex imagegen生成：第5集24镜9:16草图总览，无对白文字。</figcaption></figure>
</section>
{''.join(cards)}
</main></body></html>"""


def make_md(shots: list[Shot]) -> str:
    lines = [
        "# 第5集《活阎罗谢临舟（下）》｜导演分镜表 + Image 2.0草图提示词",
        "",
        f"- 镜头数：{len(shots)}",
        f"- 总时长：{sum(s.duration for s in shots)}秒",
        "- 草图：Image 2.0，9:16竖幅，黑白铅笔导演分镜草图",
        f"- {REF_A_LABEL}",
        f"- {REF_B_LABEL}",
        "",
    ]
    for shot in shots:
        lines += [
            f"## {shot.sid}",
            f"- 时长：{shot.duration}秒",
            f"- 景别：{shot.size}",
            f"- 拍摄视角：{shot.angle}",
            f"- 运镜：{shot.move}",
            f"- 站位：{shot.blocking}",
            f"- 画面：{shot.content}",
            f"- Image 2.0提示词：{shot.image_prompt}",
            "",
        ]
    return "\n".join(lines)


def main() -> None:
    shots = read_shots()
    write_outputs(shots)
    manifest = {
        "episode": "第5集",
        "shot_count": len(shots),
        "total_seconds": sum(s.duration for s in shots),
        "prompt_batches": [str(p.relative_to(ROOT)) for p in sorted(PROMPT_DIR.glob("batch_*.txt"))],
        "html": str((OUT_DIR / "ep5_director_storyboard_image2.html").relative_to(ROOT)),
    }
    (OUT_DIR / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
